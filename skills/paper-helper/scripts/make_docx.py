#!/usr/bin/env python3
"""把纯文本论文排版成浙江省评选规范的 Word 文档（依赖 python-docx，可选 matplotlib 出图）。

输入格式约定（UTF-8 文本文件）：
  第一行：论文标题
  「摘要：」开头行、「关键词：」开头行
  「一、二、三…」行 = 一级标题（黑体四号）
  「（一）（二）…」行 = 二级标题（黑体小四）
  「表:表题|表头1|表头2|表头3」行起始的表格块：其后每行「单元格|单元格|单元格」，空行结束（三线表，表题在上）
  「图:图片路径|图题」= 插入居中图片 + 图题（图题在下）
  「参考文献」行 = 参考文献节（悬挂缩进）

排版规范：A4；上下2.54cm 左3.0cm 右2.5cm；标题黑体小二居中；署名楷体小四居中（__待填）；
摘要/关键词楷体五号；正文宋体小四 1.5 倍行距首行缩进 2 字符；表题在上五号加粗；图题在下五号。

用法：
  python3 make_docx.py 输入.txt [输出.docx]   # 输出默认与输入同名 .docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HEI, SONG, KAI = "黑体", "宋体", "楷体"


def set_font(run, east, size, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text, east=SONG, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         indent=2.0, line=1.5, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if indent:
        pf.first_line_indent = Pt(size * indent)
    pf.line_spacing = line
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    set_font(p.add_run(text), east, size, bold)
    return p


def three_line_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = table._tbl.tblPr.makeelement(qn("w:tblBorders"), {})
    for tag, sz in (("top", "12"), ("bottom", "12")):
        borders.append(borders.makeelement(
            qn(f"w:{tag}"), {qn("w:val"): "single", qn("w:sz"): sz, qn("w:color"): "000000"}))
    for tag in ("left", "right", "insideH", "insideV"):
        borders.append(borders.makeelement(qn(f"w:{tag}"), {qn("w:val"): "none"}))
    table._tbl.tblPr.append(borders)

    def fill(cell, text, bold=False):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.2
        set_font(p.add_run(text), SONG, 10.5, bold)
        if bold:
            tc = cell._tc.get_or_add_tcPr()
            tb = tc.makeelement(qn("w:tcBorders"), {})
            tb.append(tb.makeelement(qn("w:bottom"),
                                     {qn("w:val"): "single", qn("w:sz"): "6", qn("w:color"): "000000"}))
            tc.append(tb)

    for j, h in enumerate(header):
        fill(table.rows[0].cells[j], h, bold=True)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            fill(table.rows[i].cells[j], v)


def build(txt_path: Path, out_path: Path):
    lines = txt_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)

    title = lines[0].strip() if lines else "未命名"
    para(doc, title, HEI, 18, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.3, after=6)
    para(doc, "（单位：＿＿＿＿＿＿＿＿　　作者：＿＿＿＿）", KAI, 12,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, after=8)

    in_refs = False
    i = 1
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("摘要："):
            para(doc, line, KAI, 10.5, line=1.25)
        elif line.startswith("关键词："):
            para(doc, line, KAI, 10.5, line=1.25, after=10)
        elif line == "参考文献":
            in_refs = True
            para(doc, line, HEI, 14, True, WD_ALIGN_PARAGRAPH.LEFT, 0, before=12, after=6)
        elif in_refs:
            p = para(doc, line, SONG, 10.5, WD_ALIGN_PARAGRAPH.LEFT, 0, 1.25)
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Pt(21), Pt(-21)
        elif line.startswith("表:"):
            cells = [c.strip() for c in line[2:].split("|")]
            caption, header = cells[0], cells[1:]
            rows, i = [], i + 1
            # 表格块到空行或不含 | 的行即结束（避免吞掉后续正文）
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append([c.strip() for c in lines[i].strip().split("|")])
                i += 1
            para(doc, caption, SONG, 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.2, before=6, after=4)
            three_line_table(doc, header, rows)
            i -= 1  # 回退到表格块后未消费的行（外层循环会 +1）
        elif line.startswith("图:"):
            m = [c.strip() for c in line[2:].split("|")]
            if len(m) == 2 and Path(m[0]).exists():
                doc.add_picture(m[0], width=Cm(13.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(doc, m[1] if len(m) > 1 else "", SONG, 10.5,
                 align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=2, after=8)
        elif re.match(r"^[一二三四五六七八九十]+、", line) and len(line) < 50:
            para(doc, line, HEI, 14, True, WD_ALIGN_PARAGRAPH.LEFT, 0, before=10, after=6)
        elif re.match(r"^[（(][一二三四五六]", line) and len(line) < 40:
            para(doc, line, HEI, 12, True, WD_ALIGN_PARAGRAPH.LEFT, 0, before=6, after=4)
        else:
            para(doc, line)
        i += 1

    doc.save(str(out_path))
    print(f"已生成: {out_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    build(src, dst)
