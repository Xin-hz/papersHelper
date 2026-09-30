"""文档加载与切片：支持 PDF、DOCX、DOC"""
import logging
import re
from pathlib import Path
import platform
import subprocess
import tempfile
from typing import List

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import get_settings

logger = logging.getLogger(__name__)

# 中文、字母、数字与常用标点视为有效内容字符
_MEANINGFUL = re.compile(r"[\u4e00-\u9fffA-Za-z0-9，。；：、！？…—·（）()《》<>\"'‘’“”:;,.!?\s]")
_CJK = re.compile(r"[\u4e00-\u9fff]")
_ASCII_WORD = re.compile(r"[A-Za-z]{2,}")
_XML_JUNK = re.compile(r"xmlns|xpacket|<rdf:|mwg-rs|<\?xml|<x:xmpmeta")
_CTRL_CHARS = re.compile(r"[\x00-\x08\x0b\x0e-\x1f]")
# 单文件切片上限：防止异常大文档（如内嵌大量对象的 .doc）拖垮向量化
MAX_CHUNKS_PER_FILE = 500


def _meaningful_ratio(text: str) -> float:
    if not text:
        return 0.0
    return len(_MEANINGFUL.findall(text)) / len(text)


def _is_valid_chunk(text: str) -> bool:
    """有效片段：长度够、非控制字符/乱码、且像自然语言（中文或英文散文），
    排除 .doc 内嵌图片的 XML 元数据、二进制乱码等提取残渣。"""
    t = text.strip()
    if len(t) < 50:
        return False
    if _CTRL_CHARS.search(t):
        return False
    if _meaningful_ratio(t) < 0.6:
        return False
    if len(_CJK.findall(t)) >= 10:
        return True
    if _XML_JUNK.search(t):
        return False
    return len(_ASCII_WORD.findall(t)) >= 15


def _load_doc_legacy(file_path: str) -> List[Document]:
    """加载旧版 .doc（Mac 用 textutil，其他系统请另存为 .docx）"""
    path = Path(file_path)
    if platform.system() == "Darwin":
        try:
            with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
                out_path = f.name
            subprocess.run(
                ["textutil", "-convert", "txt", "-output", out_path, str(path)],
                check=True,
                capture_output=True,
                timeout=30,
            )
            text = Path(out_path).read_text(encoding="utf-8", errors="replace")
            Path(out_path).unlink(missing_ok=True)
            # .doc 提取的文本可能混入 NUL 等控制字符，psycopg2 写库会报错
            text = text.replace("\x00", "")
        except (subprocess.CalledProcessError, FileNotFoundError, OSError) as e:
            raise ValueError(f".doc 转换失败: {e}") from e
    else:
        raise ValueError(".doc 仅在 macOS 下支持，请将 .doc 另存为 .docx 后上传。")
    if not text.strip():
        return []
    return [Document(page_content=text.strip(), metadata={"source": path.name})]


def _load_docx(file_path: str) -> List[Document]:
    """使用 python-docx 加载 DOCX"""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ImportError("请安装 python-docx: pip install python-docx")
    doc = DocxDocument(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    text = "\n\n".join(parts)
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": Path(file_path).name})]


def load_document(file_path: str) -> List[Document]:
    """根据后缀加载单个文件为 Document 列表（未切片）"""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {file_path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        loader = PyPDFLoader(str(path))
        pages = loader.load()
        return pages
    if suffix == ".docx":
        return _load_docx(str(path))
    if suffix == ".doc":
        return _load_doc_legacy(str(path))
    raise ValueError(f"不支持的文件格式: {suffix}，仅支持 .pdf / .docx / .doc")


def chunk_documents(documents: List[Document]) -> List[Document]:
    """对文档列表做递归切片，过滤过短的无效片段"""
    settings = get_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", "。", "；", " ", ""],
        length_function=len,
    )
    chunks = splitter.split_documents(documents)
    # 过滤无效片段：过短的、乱码、控制字符、XML 元数据残渣
    valid = [c for c in chunks if _is_valid_chunk(c.page_content)]
    dropped = len(chunks) - len(valid)
    if dropped:
        logger.info("已过滤 %d 个无效/乱码片段", dropped)
    if len(valid) > MAX_CHUNKS_PER_FILE:
        logger.warning("切片数 %d 超过上限 %d，已截断（文档可能异常，建议另存为 .docx 重新上传）", len(valid), MAX_CHUNKS_PER_FILE)
        valid = valid[:MAX_CHUNKS_PER_FILE]
    return valid


def load_and_chunk_file(file_path: str) -> List[Document]:
    """加载并切片单个文件，返回 Document 列表"""
    docs = load_document(file_path)
    return chunk_documents(docs)


def load_and_chunk_directory(dir_path: str) -> List[Document]:
    """加载目录下所有 PDF/DOCX 并切片"""
    allowed = {".pdf", ".docx", ".doc"}
    all_docs: List[Document] = []
    for p in Path(dir_path).rglob("*"):
        if p.is_file() and p.suffix.lower() in allowed:
            try:
                docs = load_and_chunk_file(str(p))
                for d in docs:
                    d.metadata["source"] = str(p.name)
                all_docs.extend(docs)
            except Exception as e:
                print(f"跳过文件 {p}: {e}")
    return all_docs
